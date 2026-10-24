"""
Report Exporter
Generates PDF and PPTX versions of the sales report from a sales CSV.
"""
import pandas as pd
from pathlib import Path
from datetime import datetime


def _load(csv_path: str) -> dict:
    df = pd.read_csv(csv_path, parse_dates=["date"])
    df["total"]    = pd.to_numeric(df["total"],    errors="coerce").fillna(0)
    df["quantity"] = pd.to_numeric(df["quantity"], errors="coerce").fillna(0)

    total_rev    = float(df["total"].sum())
    total_orders = len(df)

    sp = (df.groupby("salesperson")
            .agg(orders=("total", "count"), revenue=("total", "sum"))
            .sort_values("revenue", ascending=False).reset_index())

    prod = (df.groupby("product")
              .agg(units=("quantity", "sum"), revenue=("total", "sum"))
              .sort_values("revenue", ascending=False).reset_index())

    df["month"] = df["date"].dt.to_period("M")
    monthly = (df.groupby("month")
                 .agg(orders=("total", "count"), revenue=("total", "sum"))
                 .reset_index())
    monthly["month"] = monthly["month"].astype(str)

    return dict(df=df, total_rev=total_rev, total_orders=total_orders,
                sp=sp, prod=prod, monthly=monthly,
                top_person=sp.iloc[0]["salesperson"],
                best_product=prod.iloc[0]["product"],
                avg_order=total_rev / max(total_orders, 1))


# ─── PDF ─────────────────────────────────────────────────────────────────────

def generate_pdf(csv_path: str, output_path: str) -> str:
    from fpdf import FPDF, XPos, YPos

    d = _load(csv_path)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    class PDF(FPDF):
        def footer(self):
            self.set_y(-13)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(148, 163, 184)
            self.cell(0, 5, f"BusinessAutoKit by Vinesh  ·  Page {self.page_no()}",
                      align="C")

    pdf = PDF(orientation="P", format="A4")
    pdf.set_margins(14, 14, 14)
    pdf.set_auto_page_break(auto=True, margin=18)

    # ── Page 1 ───────────────────────────────────────────────────
    pdf.add_page()

    # Header bar
    pdf.set_fill_color(79, 70, 229)
    pdf.rect(0, 0, 210, 36, "F")
    pdf.set_y(8)
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 11, "Sales Performance Report", align="C",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(199, 210, 254)
    period = f"{d['monthly']['month'].iloc[0]}  to  {d['monthly']['month'].iloc[-1]}"
    pdf.cell(0, 6, f"{period}   ·   Generated {datetime.now().strftime('%d %b %Y')}",
             align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_text_color(30, 41, 59)

    # Stat boxes
    def stat_box(x, y, w, h, label, value, sub=""):
        pdf.set_fill_color(238, 242, 255)
        pdf.set_draw_color(196, 181, 253)
        pdf.rect(x, y, w, h, "FD")
        pdf.set_xy(x + 3, y + 3)
        pdf.set_font("Helvetica", "", 7)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(w - 6, 4, label.upper())
        pdf.set_xy(x + 3, y + 9)
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(30, 41, 59)
        pdf.cell(w - 6, 7, str(value))
        if sub:
            pdf.set_xy(x + 3, y + 18)
            pdf.set_font("Helvetica", "", 7)
            pdf.set_text_color(148, 163, 184)
            pdf.cell(w - 6, 4, str(sub))

    bw = 85
    stat_box(14, 42, bw, 27, "Total Revenue",
             f"Rs {d['total_rev']:,.0f}", f"{d['total_orders']} orders")
    stat_box(14 + bw + 4, 42, bw, 27, "Top Salesperson",
             d["top_person"], f"Rs {float(d['sp'].iloc[0]['revenue']):,.0f}")
    stat_box(14, 74, bw, 27, "Best Product",
             d["best_product"], f"Rs {float(d['prod'].iloc[0]['revenue']):,.0f}")
    stat_box(14 + bw + 4, 74, bw, 27, "Avg Order Value",
             f"Rs {d['avg_order']:,.0f}", "per transaction")

    def draw_table(y, title, headers, col_ws, rows_data):
        pdf.set_y(y)
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(30, 41, 59)
        pdf.cell(0, 7, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        # Header
        pdf.set_fill_color(79, 70, 229)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 8)
        for h, w in zip(headers, col_ws):
            pdf.cell(w, 7, h, fill=True)
        pdf.ln(7)
        # Rows
        for i, row_vals in enumerate(rows_data):
            even = i % 2 == 0
            pdf.set_fill_color(255, 255, 255) if even else pdf.set_fill_color(248, 250, 252)
            pdf.set_text_color(30, 41, 59)
            pdf.set_font("Helvetica", "", 8)
            for v, w in zip(row_vals, col_ws):
                pdf.cell(w, 6.5, str(v)[:28], fill=True)
            pdf.ln(6.5)

    medals = ["1st", "2nd", "3rd"]
    sp_rows = [
        [medals[i] if i < 3 else str(i + 1),
         row["salesperson"],
         str(int(row["orders"])),
         f"Rs {float(row['revenue']):,.0f}",
         f"Rs {float(row['revenue'])/max(int(row['orders']),1):,.0f}"]
        for i, row in d["sp"].iterrows()
    ]
    draw_table(108, "Salesperson Leaderboard",
               ["Rank", "Name", "Orders", "Revenue", "Avg/Order"],
               [16, 64, 18, 44, 40], sp_rows)

    # ── Page 2 ───────────────────────────────────────────────────
    pdf.add_page()

    prod_rows = [
        [row["product"], str(int(row["units"])),
         f"Rs {float(row['revenue']):,.0f}",
         f"{float(row['revenue'])/d['total_rev']*100:.1f}%"]
        for _, row in d["prod"].iterrows()
    ]
    draw_table(15, "Product Breakdown",
               ["Product", "Units Sold", "Revenue", "Share"],
               [74, 26, 46, 36], prod_rows)

    monthly_rows = [
        [str(row["month"]), str(int(row["orders"])),
         f"Rs {float(row['revenue']):,.0f}"]
        for _, row in d["monthly"].iterrows()
    ]
    draw_table(pdf.get_y() + 8, "Monthly Revenue Trend",
               ["Month", "Orders", "Revenue"],
               [52, 30, 52], monthly_rows)

    pdf.output(output_path)
    return output_path


# ─── PPTX ────────────────────────────────────────────────────────────────────

def generate_pptx(csv_path: str, output_path: str) -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    d = _load(csv_path)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    INDIGO = RGBColor(79, 70, 229)
    WHITE  = RGBColor(255, 255, 255)
    DARK   = RGBColor(30, 41, 59)
    MUTED  = RGBColor(100, 116, 139)
    LIGHT_BG = RGBColor(238, 242, 255)
    ROW_ALT  = RGBColor(248, 250, 252)

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def rect(slide, l, t, w, h, fill, line=None):
        s = slide.shapes.add_shape(1, l, t, w, h)
        s.fill.solid(); s.fill.fore_color.rgb = fill
        if line: s.line.color.rgb = line; s.line.width = Pt(0.5)
        else:     s.line.fill.background()
        return s

    def txt(slide, text, l, t, w, h, size=11, bold=False,
            color=DARK, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(l, t, w, h)
        tf = tb.text_frame; tf.word_wrap = False
        p  = tf.paragraphs[0]; p.alignment = align
        r  = p.add_run(); r.text = str(text)
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color

    def add_table_slide(title, headers, col_ws, rows_data):
        sl = prs.slides.add_slide(blank)
        rect(sl, 0, 0, prs.slide_width, Inches(1.0), INDIGO)
        txt(sl, title, Inches(0.5), Inches(0.18), Inches(12), Inches(0.7),
            size=24, bold=True, color=WHITE)
        txt(sl, "BusinessAutoKit by Vinesh",
            Inches(0.5), Inches(6.95), Inches(12), Inches(0.35),
            size=8, color=MUTED, align=PP_ALIGN.CENTER)

        ncols = len(headers)
        nrows = len(rows_data)
        total_w = sum(col_ws)
        tbl_h   = min(Inches(0.38 * (nrows + 1)), Inches(5.8))

        shape = sl.shapes.add_table(nrows + 1, ncols,
                                    Inches(0.5), Inches(1.2), total_w, tbl_h)
        tbl = shape.table

        for j, w in enumerate(col_ws):
            tbl.columns[j].width = w

        def _cell(r, c, text, bg, bold=False, sz=9):
            cell = tbl.cell(r, c)
            cell.text = str(text)[:35]
            cell.fill.solid(); cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].font.size = Pt(sz)
                p.runs[0].font.bold = bold
                p.runs[0].font.color.rgb = WHITE if bg == INDIGO else DARK

        for j, h in enumerate(headers):
            _cell(0, j, h, INDIGO, bold=True, sz=9)

        for i, row in enumerate(rows_data):
            bg = ROW_ALT if i % 2 == 1 else WHITE
            for j, v in enumerate(row):
                _cell(i + 1, j, v, bg)

        return sl

    # ── Slide 1: Title ─────────────────────────────────────────
    s1 = prs.slides.add_slide(blank)
    rect(s1, 0, 0, prs.slide_width, Inches(2.6), INDIGO)
    txt(s1, "Sales Performance Report",
        Inches(0.5), Inches(0.5), Inches(12.3), Inches(1.1),
        size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    period = f"{d['monthly']['month'].iloc[0]} – {d['monthly']['month'].iloc[-1]}"
    txt(s1, f"Period: {period}   ·   Generated {datetime.now().strftime('%d %b %Y')}   ·   BusinessAutoKit by Vinesh",
        Inches(0.5), Inches(1.65), Inches(12.3), Inches(0.5),
        size=10, color=RGBColor(199, 210, 254), align=PP_ALIGN.CENTER)

    metrics = [
        ("TOTAL REVENUE",   f"Rs {d['total_rev']:,.0f}",
         f"{d['total_orders']} orders"),
        ("TOP SALESPERSON", d["top_person"],
         f"Rs {float(d['sp'].iloc[0]['revenue']):,.0f}"),
        ("BEST PRODUCT",    d["best_product"],
         f"Rs {float(d['prod'].iloc[0]['revenue']):,.0f}"),
        ("AVG ORDER VALUE", f"Rs {d['avg_order']:,.0f}", "per transaction"),
    ]

    bw, bh = Inches(2.9), Inches(1.85)
    gap    = Inches(0.18)
    sx, sy = Inches(0.53), Inches(2.85)

    for i, (label, value, sub) in enumerate(metrics):
        x = sx + i * (bw + gap)
        rect(s1, x, sy, bw, bh, LIGHT_BG, RGBColor(196, 181, 253))
        txt(s1, label, x + Inches(.15), sy + Inches(.15), bw, Inches(.3),
            size=7, color=MUTED)
        txt(s1, value, x + Inches(.15), sy + Inches(.45), bw, Inches(.75),
            size=17, bold=True, color=DARK)
        txt(s1, sub, x + Inches(.15), sy + Inches(1.3), bw, Inches(.35),
            size=9, color=MUTED)

    txt(s1, "Built by Vinesh  ·  Python Automation  ·  AWS Certified  ·  Fiverr Portfolio",
        Inches(.5), Inches(6.95), Inches(12.3), Inches(.35),
        size=9, color=MUTED, align=PP_ALIGN.CENTER)

    # ── Slide 2: Leaderboard ────────────────────────────────────
    medals = ["🥇 1st", "🥈 2nd", "🥉 3rd"]
    sp_rows = [
        [medals[i] if i < 3 else str(i + 1),
         row["salesperson"], str(int(row["orders"])),
         f"Rs {float(row['revenue']):,.0f}",
         f"Rs {float(row['revenue'])/max(int(row['orders']),1):,.0f}"]
        for i, row in d["sp"].iterrows()
    ]
    add_table_slide("Salesperson Leaderboard",
                    ["Rank", "Name", "Orders", "Total Revenue", "Avg Per Order"],
                    [Inches(1.0), Inches(4.0), Inches(1.3), Inches(3.0), Inches(3.0)],
                    sp_rows)

    # ── Slide 3: Products ───────────────────────────────────────
    prod_rows = [
        [row["product"], str(int(row["units"])),
         f"Rs {float(row['revenue']):,.0f}",
         f"{float(row['revenue'])/d['total_rev']*100:.1f}%"]
        for _, row in d["prod"].iterrows()
    ]
    add_table_slide("Product Breakdown",
                    ["Product", "Units Sold", "Revenue", "Share of Total"],
                    [Inches(5.0), Inches(2.0), Inches(3.0), Inches(2.3)],
                    prod_rows)

    # ── Slide 4: Monthly ────────────────────────────────────────
    mrows = []
    prev  = None
    for _, row in d["monthly"].iterrows():
        change = ""
        if prev is not None:
            pct = (float(row["revenue"]) - prev) / prev * 100
            change = f"+{pct:.1f}%" if pct >= 0 else f"{pct:.1f}%"
        mrows.append([str(row["month"]), str(int(row["orders"])),
                      f"Rs {float(row['revenue']):,.0f}", change])
        prev = float(row["revenue"])

    add_table_slide("Monthly Revenue Trend",
                    ["Month", "Orders", "Revenue", "vs Previous"],
                    [Inches(3.0), Inches(2.0), Inches(3.5), Inches(2.0)],
                    mrows)

    prs.save(output_path)
    return output_path
